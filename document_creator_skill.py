import logging
import os
from datetime import datetime, timezone
from typing import ClassVar

from base import Skill
from config import (
    DOCS_OUTPUT_DIR,
    OPENAI_API_KEY,
    OPENAI_BASE_URL,
    OPENAI_MODEL,
)

try:
    from docx import Document  # pyright: ignore[reportMissingImports]
except ImportError:
    Document = None
try:
    from openai import OpenAI  # pyright: ignore[reportMissingImports]
except ImportError:
    OpenAI = None

logger = logging.getLogger(__name__)

class DocumentCreatorSkill(Skill):
    name = "document_creator"
    keywords: ClassVar[list[str]] = [
        "create a document", "make a document", "write a document",
        "create a word doc", "make a word doc", "generate a report",
    ]

    _TRIGGERS: ClassVar[list[str]] = [
        "create a document about", "create a document with",
        "make a document about", "make a document with",
        "write a document about", "write a document with",
        "create a report on", "create a report about",
    ]

    def __init__(self):
        if not OPENAI_API_KEY:
            logger.warning("DocumentCreatorSkill: OPENAI_API_KEY is not set.")
            self.client = None
            return
        self.client = OpenAI(api_key=OPENAI_API_KEY, base_url=OPENAI_BASE_URL)

    def handle(self, user_input: str) -> str:
        if self.client is None:
            return "Document creation is unavailable - Groq key missing."

        try:
            os.makedirs(DOCS_OUTPUT_DIR, exist_ok=True)
            content_request = self._strip_trigger(user_input)

            response = self.client.chat.completions.create(
                model=OPENAI_MODEL,
                messages=[
                    {
                        "role": "system",
                        "content": (
                            "You create document outlines. "
                            "Format strictly as: TITLE: <title>\\n---\\n<body>"
                        ),
                    },
                    {
                        "role": "user",
                        "content": f"Create a document about: {content_request}",
                    },
                ],
                temperature=0.5,
                max_tokens=500,
            )
            structured = response.choices[0].message.content.strip()
            title, body = self._parse_response(structured, content_request)

            safe_title = "".join(
                c for c in title if c.isalnum() or c in ("-", "_")
            ).strip()[:50] or "document"
            timestamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
            filepath = os.path.join(
                DOCS_OUTPUT_DIR, f"{safe_title}_{timestamp}.docx"
            )

            doc = Document()
            doc.add_heading(title, 0)
            for paragraph in body.split("\n\n"):
                if paragraph.strip():
                    doc.add_paragraph(paragraph.strip())
            doc.save(filepath)

            return f"Done. Created '{title}' in your documents folder."

        except (
            OSError,
            ValueError,
            TypeError,
            AttributeError,
            RuntimeError,
        ) as e:
            logger.error(f"Document creation error: {e}")
            return "Sorry, I had trouble creating that document."

    def _strip_trigger(self, text: str) -> str:
        lower = text.lower()
        for trigger in self._TRIGGERS:
            if lower.startswith(trigger):
                return text[len(trigger):].strip()
        return text.strip()

    def _parse_response(self, response: str, fallback_title: str) -> tuple:
        lines = response.splitlines()
        title = fallback_title.strip() or "Document"
        body = ""

        for i, line in enumerate(lines):
            if line.strip().startswith("TITLE:"):
                title = line.split(":", 1)[1].strip() or title
                if i + 1 < len(lines) and lines[i + 1].strip() == "---":
                    body = "\n".join(lines[i + 2:]).strip()
                else:
                    body = "\n".join(lines[i + 1:]).strip()
                break
        else:
            body = "\n".join(lines).strip()

        return title, body

